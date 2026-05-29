"""
PUA (Praktiniu uzduociu ataskaita) sujungimo skriptas.

Sujungia 7 atskirus PU .docx failus i viena PUA dokumenta pagal
Vilniaus universiteto Kauno fakulteto Informatikos inzinerijos krypties
akademiniu rasto darbu metodinius nurodymus.

Naudojimas:
    python merge_pua.py

Reikalavimai:
    pip install python-docx docxcompose

Autorius: Andrius Vargonas
Data: 2026-05
"""

import os
from copy import deepcopy

from docx import Document
from docx.shared import Pt, Mm, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docxcompose.composer import Composer


# ============================================================
# KONFIGURACIJA
# ============================================================

REPO_DIR = "/projects/sandbox/a"
OUT_DIR = os.path.join(REPO_DIR, "PUA_FINAL")
OUT_FILE = os.path.join(OUT_DIR, "PUA_Andrius_Vargonas.docx")

# (failo_pavadinimas, skyriaus_numeris, skyriaus_antraste)
PU_FILES = [
    ("PU1_Andrius_Vargonas.docx",  1, "PRAKTINĖ UŽDUOTIS PU1"),
    ("PU2_Andrius_Vargonas.docx",  2, "PRAKTINĖ UŽDUOTIS PU2"),
    ("PU3_Andrius_Vargonas1.docx", 3, "PRAKTINĖ UŽDUOTIS PU3"),
    ("PU4_Andrius_Vargonas.docx",  4, "PRAKTINĖ UŽDUOTIS PU4"),
    ("PU5_Andrius_Vargonas.docx",  5, "PRAKTINĖ UŽDUOTIS PU5"),
    ("PU6_Andrius_Vargonas.docx",  6, "PRAKTINĖ UŽDUOTIS PU6"),
    ("PU7_Andrius_Vargonas.docx",  7, "PRAKTINĖ UŽDUOTIS PU7"),
]

# Lietuviskos antrastes su diakritikais - naudojamos dokumento tekste
SKYRIAU_ANTRASTES_LT = {
    1: "PRAKTINĖ UŽDUOTIS PU1",
    2: "PRAKTINĖ UŽDUOTIS PU2",
    3: "PRAKTINĖ UŽDUOTIS PU3",
    4: "PRAKTINĖ UŽDUOTIS PU4",
    5: "PRAKTINĖ UŽDUOTIS PU5",
    6: "PRAKTINĖ UŽDUOTIS PU6",
    7: "PRAKTINĖ UŽDUOTIS PU7",
}


# ============================================================
# PAGALBINES FUNKCIJOS - Word OOXML laukai (fields)
# ============================================================

def _add_field(paragraph, instr_text, placeholder="(atnaujinkite F9)"):
    """Iterpia Word lauka (field) su nurodyta instrukcija (pvz. TOC, PAGE)."""
    run = paragraph.add_run()
    r = run._r

    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    r.append(fld_begin)

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instr_text
    r.append(instr)

    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    r.append(fld_sep)

    t = OxmlElement('w:t')
    t.text = placeholder
    r.append(t)

    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    r.append(fld_end)

    return run


def _set_run_font(run, name="Times New Roman", size_pt=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    rFonts.set(qn('w:eastAsia'), name)


def _add_page_break(paragraph):
    run = paragraph.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)


# ============================================================
# DOKUMENTO NUSTATYMAI
# ============================================================

def configure_section(section):
    """A4 lapas, paraštes ir header/footer atstumai pagal VU KnF reikalavimus."""
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(12.5)
    section.different_first_page_header_footer = True


def configure_default_styles(doc):
    """Times New Roman 12pt, 1.5 line spacing, justify, 1.25cm pirma eilutes itrauka."""
    # Normal stilius
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)

    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)

    # Heading 1 stilius - skyriu antrastes (didziosiomis, 14 pt, center, bold)
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    rPr = h1.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    pf = h1.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(12)
    pf.line_spacing = 1.5
    pf.page_break_before = True  # skyrius pradedamas naujame lape

    # Caption stilius (paveikslams ir lentelems) - 11 pt, bold
    if 'Caption' in [s.name for s in doc.styles]:
        cap = doc.styles['Caption']
        cap.font.name = 'Times New Roman'
        cap.font.size = Pt(11)
        cap.font.bold = True
        cap.font.italic = False
        cap.font.color.rgb = RGBColor(0, 0, 0)


def setup_footer_page_numbers(section):
    """Idek puslapio numeri desineje footer'yje. Pirmasis puslapis (titulinis) - tuscias."""
    # Pagrindinis footer - rodo puslapio numeri desineje
    footer = section.footer
    if footer.paragraphs:
        para = footer.paragraphs[0]
    else:
        para = footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = para.paragraph_format
    pf.first_line_indent = Cm(0)
    # PAGE field
    _add_field(para, " PAGE   \\* MERGEFORMAT ", placeholder="2")

    # Pirmojo puslapio footer - tuscias (titulinis nenumeruojamas)
    first_footer = section.first_page_footer
    if not first_footer.paragraphs:
        first_footer.add_paragraph()
    fp = first_footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT


# ============================================================
# TITULINIS LAPAS
# ============================================================

def add_title_page(doc):
    """Titulinis lapas pagal VU KnF reikalavimus."""
    def _para(text, *, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER,
              space_before=0, space_after=0):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.5
        run = p.add_run(text)
        _set_run_font(run, size_pt=size, bold=bold)
        return p

    # Virsus - universitetas
    _para("VILNIAUS UNIVERSITETAS", size=14, bold=True, space_after=6)
    _para("KAUNO FAKULTETAS", size=14, bold=True, space_after=6)
    _para("SOCIALINIŲ MOKSLŲ IR TAIKOMOSIOS INFORMATIKOS INSTITUTAS",
          size=14, bold=True, space_after=24)

    # Studiju programa ir autorius
    _para("Studijų programa: Marketingo technologijos", size=12, space_after=120)

    _para("Andrius Vargonas", size=14, bold=True, space_after=120)

    # Pavadinimas
    _para("PRAKTINIŲ UŽDUOČIŲ ATASKAITA", size=18, bold=True, space_after=6)
    _para("(PUA)", size=18, bold=True, space_after=24)

    _para("Dalykas: Informacijos sistemos ir duomenų bazės",
          size=12, space_after=240)

    # Apacia - vieta ir metai
    _para("Kaunas", size=12, space_after=0)
    _para("2026", size=12, space_after=0)

    # Page break i kita lapa
    last_p = doc.paragraphs[-1]
    _add_page_break(last_p)


# ============================================================
# TURINYS, PAVEIKSLU IR LENTELIU SARASAI
# ============================================================

def add_section_heading_no_toc(doc, text):
    """Skyriaus tipo antraste (TURINYS, PAVEIKSLU SARASAS, LENTELIU SARASAS),
    kuri NETURI buti automatiniame TOC. Stilius - Normal su rankiniu formatavimu."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(18)
    pf.line_spacing = 1.5
    run = p.add_run(text)
    _set_run_font(run, size_pt=14, bold=True)
    return p


def add_toc(doc, title_text, toc_instr):
    """Prideda antraste + TOC field."""
    add_section_heading_no_toc(doc, title_text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    _add_field(p, toc_instr, placeholder="Po atidarymo paspauskite Ctrl+A, F9 turiniui atnaujinti.")
    # Page break po TOC
    last = doc.paragraphs[-1]
    _add_page_break(last)


# ============================================================
# SKYRIU ANTRASTES PRIES KIEKVIENA PU
# ============================================================

def add_chapter_heading(doc, number, title):
    """Heading 1: '1. PRAKTINE UZDUOTIS PU1' - bus surinktas i TURINYS automatiskai."""
    text = "{}. {}".format(number, title)
    p = doc.add_heading(text, level=1)
    # Heading 1 jau yra page_break_before per stilius
    return p


# ============================================================
# DOKUMENTO PRIJUNGIMAS PER docxcompose
# ============================================================

def append_pu_document(composer, pu_path):
    """Prijungia visa PU .docx faila prie master."""
    sub_doc = Document(pu_path)
    composer.append(sub_doc)


# ============================================================
# PAGRINDINE FUNKCIJA
# ============================================================

def build_pua():
    print("=" * 60)
    print("PUA generavimas pradedamas")
    print("=" * 60)

    # 1. Sukurk MASTER dokumenta
    master = Document()
    configure_default_styles(master)

    # Pirmoji (ir vienintele master) sekcija
    section = master.sections[0]
    configure_section(section)
    setup_footer_page_numbers(section)

    print("[1/5] Sukurtas master dokumentas su VU formatavimu")

    # 2. Titulinis lapas
    add_title_page(master)
    print("[2/5] Pridetas titulinis lapas")

    # 3. TURINYS (su Word TOC field)
    add_toc(master, "TURINYS", 'TOC \\o "1-3" \\h \\z \\u')

    # 4. PAVEIKSLU SARASAS
    add_toc(master, "PAVEIKSLŲ SĄRAŠAS", 'TOC \\h \\z \\c "Pav"')

    # 5. LENTELIU SARASAS
    add_toc(master, "LENTELIŲ SĄRAŠAS", 'TOC \\h \\z \\c "Lentelė"')

    print("[3/5] Prideti TURINYS, PAVEIKSLU SARASAS, LENTELIU SARASAS")

    # Issaugok master, kad jis butu paruostas docxcompose'ui
    tmp_master = os.path.join(OUT_DIR, "_master_tmp.docx")
    master.save(tmp_master)

    # 6. Atidarom is naujo per Composer ir prijungiam PU dokumentus
    base = Document(tmp_master)
    composer = Composer(base)

    print("[4/5] Pradedamas PU dokumentu prijungimas...")
    for filename, number, title in PU_FILES:
        pu_path = os.path.join(REPO_DIR, filename)
        if not os.path.exists(pu_path):
            raise FileNotFoundError("Nerastas PU failas: {}".format(pu_path))

        # Pridek skyriaus antraste prie base PRIES atvera PU
        add_chapter_heading(base, number, SKYRIAU_ANTRASTES_LT[number])

        # Tada Composer prijungia visa PU dokumenta
        sub = Document(pu_path)
        composer.append(sub)

        size_kb = os.path.getsize(pu_path) // 1024
        print("    + {} (skyrius {}, {} KB)".format(filename, number, size_kb))

    # 7. Issaugom galutini PUA dokumenta
    composer.save(OUT_FILE)

    # Pasaliname laikina
    if os.path.exists(tmp_master):
        os.remove(tmp_master)

    final_size_kb = os.path.getsize(OUT_FILE) // 1024
    print("[5/5] PUA issaugotas: {} ({} KB)".format(OUT_FILE, final_size_kb))
    print("=" * 60)
    print("ATIDARE WORD'E AR LIBREOFFICE'E:")
    print("  1. Ctrl+A (paryskinti viska)")
    print("  2. F9 (atnaujinti turini ir laukus)")
    print("  3. Patikrinti puslapiu numeravima ir TOC'us")
    print("=" * 60)


if __name__ == "__main__":
    build_pua()
