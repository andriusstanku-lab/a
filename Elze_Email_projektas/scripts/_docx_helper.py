"""Bendras helper'is VU stiliaus .docx dokumentams kurti.

Stilius: Times New Roman 12pt, 1.5 tarpai, pateisinta lygiavimas,
puslapių numeracija footeryje. TURINYS – statinis tekstas, ne lauko kodas.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _add_page_number(paragraph):
    """Įdeda automatinį puslapio numerį (PAGE field) paragrafe."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def new_document():
    """Sukuria naują dokumentą su VU stiliaus formatavimu."""
    doc = Document()

    # Bazinis stilius
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    # Lithuanian font fallback
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')

    # Paraščių nustatymas (VU standartas: viršus/apačia 2cm, kairė 3cm, dešinė 1.5cm)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    # Puslapio numeris footeryje
    section = doc.sections[0]
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(fp)

    return doc


def add_title(doc, title):
    """Pagrindinė antraštė – centruota, Bold, 14pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = 'Times New Roman'


def add_h1(doc, text):
    """1 lygio paantraštė – Bold, 13pt, kairėn."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.name = 'Times New Roman'


def add_h2(doc, text):
    """2 lygio paantraštė – Bold, 12pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = 'Times New Roman'


def add_para(doc, text, bold=False, italic=False, align='justify'):
    """Įprasta pastraipa, justify, 1.5 tarpai, įtrauka 1cm."""
    p = doc.add_paragraph()
    if align == 'justify':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'left':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(1.0)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = bold
    r.italic = italic
    return p


def add_bullet(doc, text, bold_prefix=None):
    """Sąrašo punktas su • prefix'u (be auto-numeravimo, kad atrodytų švariai)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run('•  ')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    if bold_prefix:
        r2 = p.add_run(bold_prefix)
        r2.bold = True
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
        r3 = p.add_run(text)
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(12)
    else:
        r2 = p.add_run(text)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
    return p


def add_numbered(doc, num, text, bold_prefix=None):
    """Numeruotas punktas (rankiniu būdu suformuotas)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    r = p.add_run(f'{num}.  ')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    if bold_prefix:
        r2 = p.add_run(bold_prefix)
        r2.bold = True
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
        r3 = p.add_run(text)
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(12)
    else:
        r2 = p.add_run(text)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
    return p


def add_static_toc(doc, entries):
    """STATINIS turinys – paprastas tekstas, NE lauko kodas.

    entries: list of tuples (sektion_pavadinimas, puslapio_numeris)
    """
    add_h1(doc, 'TURINYS')
    for name, page in entries:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        # Pavadinimas + taškeliai + puslapis (vizualiai gražu)
        dots_count = max(3, 70 - len(name) - len(str(page)))
        line = f'{name} {"."*dots_count} {page}'
        r = p.add_run(line)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
    # Tarpas po turinio
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)


def add_table_simple(doc, headers, rows, col_widths_cm=None):
    """Paprasta lentelė su antrašte ir eilutėmis.

    headers: [str, ...] – antraštės
    rows: [[str, ...], ...] – eilutės
    col_widths_cm: list of float, neprivaloma
    """
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True

    # Antraštė
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = 'Times New Roman'
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Eilutės
    for ri, row in enumerate(rows, start=1):
        cells = table.rows[ri].cells
        for ci, val in enumerate(row):
            cells[ci].text = ''
            p = cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.size = Pt(11)
            r.font.name = 'Times New Roman'
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Stulpelių pločiai
    if col_widths_cm:
        for ri in range(len(table.rows)):
            for ci, w in enumerate(col_widths_cm):
                table.rows[ri].cells[ci].width = Cm(w)

    # Tarpas po lentelės
    doc.add_paragraph()
    return table


def add_quote_box(doc, text):
    """Citatos / svarbios pastraipos blokas (Italic, kursyvinis, įtrauktas)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.right_indent = Cm(1.0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)


def add_pagebreak(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    from docx.enum.text import WD_BREAK
    r.add_break(WD_BREAK.PAGE)


def add_header_block(doc, university, faculty, course, project_title, group, channel, date):
    """Pirmojo puslapio antraštės blokas (vietoje pilno titulinio)."""
    # Universitetas
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(university)
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = 'Times New Roman'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(faculty)
    r.font.size = Pt(11)
    r.font.name = 'Times New Roman'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(course)
    r.italic = True
    r.font.size = Pt(11)
    r.font.name = 'Times New Roman'

    # Tarpas
    doc.add_paragraph()

    # Pavadinimas
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(project_title)
    r.bold = True
    r.font.size = Pt(15)
    r.font.name = 'Times New Roman'

    # Kanalas / poprojektis
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(f'Pasirinktas kanalas: {channel}')
    r.italic = True
    r.font.size = Pt(12)
    r.font.name = 'Times New Roman'

    # Grupė ir data
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f'Grupė: {group}')
    r.font.size = Pt(11)
    r.font.name = 'Times New Roman'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(f'Data: {date}')
    r.font.size = Pt(11)
    r.font.name = 'Times New Roman'
