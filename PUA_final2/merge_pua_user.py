#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PUA (Praktinių užduočių ataskaita) sujungimas iš vartotojo 7 .docx failų.
Autorius: Andrius Vargonas, VU Kauno fakultetas
"""
import os
import sys
from copy import deepcopy

from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from docxcompose.composer import Composer

PUA_DIR = "/projects/sandbox/PUA_merge"
OUTPUT = "PUA_Andrius_Vargonas.docx"

PU_FILES = [
    ("PU1", "PU1_Andrius_Vargonas.docx", "Informacinės sistemos ir duomenų bazės"),
    ("PU2", "PU2_Andrius_Vargonas.docx", "Reliacinis duomenų modelis"),
    ("PU3", "PU3_Andrius_Vargonas1.docx", "Microsoft Access aplinka"),
    ("PU4", "PU4_Andrius_Vargonas.docx", "Duomenų bazės kūrimas Java priemonėmis"),
    ("PU5", "PU5_Andrius_Vargonas.docx", "Duomenų bazių pavyzdžiai"),
    ("PU6", "PU6_Andrius_Vargonas.docx", "Normalizavimas iki 3NF"),
    ("PU7", "PU7_Andrius_Vargonas.docx", "Vaizdai (Views) ir užklausos"),
]


# -------------------- Pagalbinės funkcijos --------------------
def _set_page_setup(section):
    """A4, paraštės: viršus 20, apačia 20, kairė 25, dešinė 15 mm."""
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)


def _add_field(paragraph, instr_text):
    """Įterpia Word lauką (TOC, PAGE, ir pan.) į pastraipą."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    fldChar1.set(qn('w:dirty'), 'true')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = instr_text

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')

    placeholder = OxmlElement('w:t')
    placeholder.text = "Atnaujinkite šį lauką (F9 / Update Field)"

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    rPlace = OxmlElement('w:r')
    rPlace.append(placeholder)
    paragraph._p.append(rPlace)
    r2 = OxmlElement('w:r')
    r2.append(fldChar3)
    paragraph._p.append(r2)


def _add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def _heading_style(doc, name, size_pt, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                   space_before=12, space_after=6):
    style = doc.styles[name] if name in doc.styles else None
    if style is None:
        return
    f = style.font
    f.name = "Times New Roman"
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5


def _set_default_font(doc):
    style = doc.styles['Normal']
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)


# -------------------- Master dokumentas --------------------
def build_master(path):
    doc = Document()
    _set_default_font(doc)
    _set_page_setup(doc.sections[0])

    # Antraštinis lapas
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VILNIAUS UNIVERSITETAS")
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("KAUNO FAKULTETAS")
    r.bold = True
    r.font.size = Pt(14)

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Andrius Vargonas")
    r.font.size = Pt(13)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PRAKTINIŲ UŽDUOČIŲ ATASKAITA (PUA)")
    r.bold = True
    r.font.size = Pt(18)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Duomenų bazės")
    r.italic = True
    r.font.size = Pt(13)

    for _ in range(10):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Kaunas, 2026")
    r.font.size = Pt(12)

    _add_page_break(doc)

    # TURINYS
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TURINYS")
    r.bold = True
    r.font.size = Pt(14)
    p_toc = doc.add_paragraph()
    _add_field(p_toc, r' TOC \o "1-3" \h \z \u ')
    _add_page_break(doc)

    # PAVEIKSLŲ SĄRAŠAS
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PAVEIKSLŲ SĄRAŠAS")
    r.bold = True
    r.font.size = Pt(14)
    p_fig = doc.add_paragraph()
    _add_field(p_fig, r' TOC \h \z \c "Pav" ')
    _add_page_break(doc)

    # LENTELIŲ SĄRAŠAS
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LENTELIŲ SĄRAŠAS")
    r.bold = True
    r.font.size = Pt(14)
    p_lent = doc.add_paragraph()
    _add_field(p_lent, r' TOC \h \z \c "Lentelė" ')
    _add_page_break(doc)

    doc.save(path)
    return path


# -------------------- Skyriaus skirtukas --------------------
def make_chapter_separator(num, code, title, path):
    """Sukuria mažą docx su skyriaus antrašte (Heading 1)."""
    doc = Document()
    _set_default_font(doc)
    _set_page_setup(doc.sections[0])
    h = doc.add_heading(level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = h.add_run(f"{num}. {code} - {title}")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.save(path)
    return path


# -------------------- MAIN --------------------
def main():
    master_path = "/tmp/master_pua.docx"
    build_master(master_path)
    print(f"[OK] Master sukurtas: {master_path}")

    composer = Composer(Document(master_path))

    missing = []
    added = []
    for idx, (code, fname, title) in enumerate(PU_FILES, start=1):
        full = os.path.join(PUA_DIR, fname)
        if not os.path.exists(full):
            print(f"[TRUKSTA] {code}: {fname}")
            missing.append(fname)
            continue

        # Skyriaus skirtukas su Heading 1 antrašte (kad TOC matytų skyrius)
        sep_path = f"/tmp/sep_{code}.docx"
        make_chapter_separator(idx, code, title, sep_path)
        composer.append(Document(sep_path))

        # Pats PU dokumento turinys (su jo tituliniu, TOC ir kt.)
        composer.append(Document(full))
        added.append(code)
        print(f"[+] Prijungta: {code} - {fname}")

    out_path = os.path.join(PUA_DIR, OUTPUT)
    composer.save(out_path)
    size = os.path.getsize(out_path)
    print(f"\n[REZULTATAS]")
    print(f"  Failas:  {out_path}")
    print(f"  Dydis:   {size:,} baitai ({size/1024:.1f} KB)")
    print(f"  Pridėta: {', '.join(added)}")
    if missing:
        print(f"  TRUKSTA: {', '.join(missing)}")


if __name__ == "__main__":
    main()
