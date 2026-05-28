"""
PU1 užduoties Word dokumento generavimas
Autorius: Andrius Vargonas
Pagal: Akademinių rašto darbų metodiniai nurodymai (VU Kauno fakultetas, Informatikos inžinerijos kryptis)
"""

from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement


# ------------- KONFIGŪRACIJA -------------
FONT_NAME = "Times New Roman"
FONT_SIZE_BODY = 12
FONT_SIZE_H1 = 14
FONT_SIZE_H2 = 12
FONT_SIZE_TABLE_TITLE = 11
FONT_SIZE_TABLE_DATA = 10
FONT_SIZE_SOURCE = 9


# ------------- PAGALBINĖS FUNKCIJOS -------------
def set_run_font(run, name=FONT_NAME, size=FONT_SIZE_BODY, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    # Ensure font also for East Asian (helps consistency)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    rFonts.set(qn('w:eastAsia'), name)


def set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         line_spacing=1.5, first_line_indent=Cm(1.25),
                         space_before=0, space_after=0):
    pf = p.paragraph_format
    pf.alignment = alignment
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)


def add_body_paragraph(doc, text, indent=True):
    """Standartinė pastraipa: 12pt, TNR, 1.5 spacing, justify, įtrauka 1.25cm"""
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.5,
        first_line_indent=Cm(1.25) if indent else None,
        space_before=0, space_after=0
    )
    run = p.add_run(text)
    set_run_font(run, size=FONT_SIZE_BODY)
    return p


def add_h1(doc, text, numbered=True, page_break_before=True):
    """Skyriaus antraštė: 14pt Bold UPPERCASE centruota, naujame lape"""
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.5,
        first_line_indent=None,
        space_before=0, space_after=12
    )
    if page_break_before:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text.upper())
    set_run_font(run, size=FONT_SIZE_H1, bold=True)
    # Style identifier for TOC (using built-in Heading 1)
    if numbered:
        p.style = doc.styles['Heading 1']
        # override style font
        for r in p.runs:
            set_run_font(r, size=FONT_SIZE_H1, bold=True)
            r.font.color.rgb = RGBColor(0, 0, 0)
    else:
        # Use Heading 1 style still for TOC inclusion but mark? Actually unnumbered chapters
        # should NOT appear in numbered TOC. We'll add them with Heading 1 style and let TOC
        # show them. The "TURINYS" itself we exclude via separate handling.
        p.style = doc.styles['Heading 1']
        for r in p.runs:
            set_run_font(r, size=FONT_SIZE_H1, bold=True)
            r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_h2(doc, text):
    """Poskyrio antraštė: 12pt Bold, kairėje, 2 eilučių tarpas prieš, 1 po"""
    # 2 line breaks before
    p_before = doc.add_paragraph()
    set_paragraph_format(p_before, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                         line_spacing=1.5, first_line_indent=None,
                         space_before=0, space_after=0)

    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.5,
        first_line_indent=None,
        space_before=12, space_after=6
    )
    run = p.add_run(text)
    set_run_font(run, size=FONT_SIZE_H2, bold=True)
    p.style = doc.styles['Heading 2']
    for r in p.runs:
        set_run_font(r, size=FONT_SIZE_H2, bold=True)
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_table_caption(doc, number_text, title_text):
    """Lentelės numeris (dešinėje) + pavadinimas (centre) 11pt Bold"""
    p_num = doc.add_paragraph()
    set_paragraph_format(p_num, alignment=WD_ALIGN_PARAGRAPH.RIGHT,
                         line_spacing=1.0, first_line_indent=None,
                         space_before=6, space_after=0)
    r1 = p_num.add_run(number_text)
    set_run_font(r1, size=FONT_SIZE_TABLE_TITLE, bold=True)

    p_title = doc.add_paragraph()
    set_paragraph_format(p_title, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         line_spacing=1.0, first_line_indent=None,
                         space_before=0, space_after=6)
    r2 = p_title.add_run(title_text)
    set_run_font(r2, size=FONT_SIZE_TABLE_TITLE, bold=True)


def add_source_note(doc, text):
    """Šaltinio užrašas po lentele/paveikslu, 9pt"""
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                         line_spacing=1.0, first_line_indent=None,
                         space_before=2, space_after=12)
    r = p.add_run(text)
    set_run_font(r, size=FONT_SIZE_SOURCE, italic=True)


def add_field(paragraph, field_code):
    """Add a Word field (e.g., PAGE, TOC) as XML."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = field_code

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    run = paragraph.add_run()
    run._element.append(fldChar1)
    run._element.append(instrText)
    run._element.append(fldChar2)
    run._element.append(fldChar3)
    return run


# ------------- PRADŽIA -------------
doc = Document()

# Numatytasis stilius — Times New Roman 12
style = doc.styles['Normal']
style.font.name = FONT_NAME
style.font.size = Pt(FONT_SIZE_BODY)
rPr = style.element.get_or_add_rPr()
rFonts = rPr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rPr.append(rFonts)
rFonts.set(qn('w:ascii'), FONT_NAME)
rFonts.set(qn('w:hAnsi'), FONT_NAME)
rFonts.set(qn('w:cs'), FONT_NAME)
rFonts.set(qn('w:eastAsia'), FONT_NAME)

# Paraščių nustatymas: viršus 20mm, apačia 20mm, kairė 25mm, dešinė 15mm
section = doc.sections[0]
section.page_height = Mm(297)
section.page_width = Mm(210)
section.top_margin = Mm(20)
section.bottom_margin = Mm(20)
section.left_margin = Mm(25)
section.right_margin = Mm(15)
section.header_distance = Mm(12.5)
section.footer_distance = Mm(12.5)

# Stiliaus nustatymas Heading 1 ir Heading 2 (pagal akademinius reikalavimus)
for style_name, sz, bold in [('Heading 1', FONT_SIZE_H1, True), ('Heading 2', FONT_SIZE_H2, True)]:
    s = doc.styles[style_name]
    s.font.name = FONT_NAME
    s.font.size = Pt(sz)
    s.font.bold = bold
    s.font.color.rgb = RGBColor(0, 0, 0)
    rPr = s.element.get_or_add_rPr()
    rFonts2 = rPr.find(qn('w:rFonts'))
    if rFonts2 is None:
        rFonts2 = OxmlElement('w:rFonts')
        rPr.append(rFonts2)
    rFonts2.set(qn('w:ascii'), FONT_NAME)
    rFonts2.set(qn('w:hAnsi'), FONT_NAME)
    rFonts2.set(qn('w:cs'), FONT_NAME)
    rFonts2.set(qn('w:eastAsia'), FONT_NAME)


# ============================================
# 1. ANTRAŠTINIS LAPAS (1 PRIEDAS pagal pavyzdį)
# ============================================
def add_title_page_line(doc, text, size, bold=False, space_after_pt=0):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         line_spacing=1.0, first_line_indent=None,
                         space_before=0, space_after=space_after_pt)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


add_title_page_line(doc, "VILNIAUS UNIVERSITETO", FONT_SIZE_H1, bold=True)
add_title_page_line(doc, "KAUNO FAKULTETAS", FONT_SIZE_H1, bold=True, space_after_pt=30)
add_title_page_line(doc, "SOCIALINIŲ MOKSLŲ IR TAIKOMOSIOS INFORMATIKOS INSTITUTAS",
                    FONT_SIZE_H1, bold=False, space_after_pt=154)
add_title_page_line(doc, "Marketingo technologijų studijų programa",
                    FONT_SIZE_BODY, bold=False, space_after_pt=90)
add_title_page_line(doc, "ANDRIUS VARGONAS", FONT_SIZE_BODY, bold=True, space_after_pt=90)
add_title_page_line(doc, "INFORMACINIŲ SISTEMŲ IR DUOMENŲ BAZIŲ KŪRIMO REIKALAVIMŲ PALYGINIMAS",
                    FONT_SIZE_H1, bold=True, space_after_pt=30)
add_title_page_line(doc, "Informacijos sistemos ir duomenų bazės",
                    FONT_SIZE_BODY, bold=False, space_after_pt=240)
add_title_page_line(doc, "Kaunas 2026", FONT_SIZE_BODY, bold=False)

# ----- Section break: titulinis lapas atskiras nuo likusio dokumento -----
new_section = doc.add_section(WD_SECTION.NEW_PAGE)
new_section.page_height = Mm(297)
new_section.page_width = Mm(210)
new_section.top_margin = Mm(20)
new_section.bottom_margin = Mm(20)
new_section.left_margin = Mm(25)
new_section.right_margin = Mm(15)
new_section.header_distance = Mm(12.5)
new_section.footer_distance = Mm(12.5)

# Atjungti "Link to Previous" antros sekcijos poraštei
new_section.footer.is_linked_to_previous = False
# Pirmosios sekcijos (titulinio) poraštė liks tuščia
doc.sections[0].footer.is_linked_to_previous = False

# Pridėti puslapio numerį dešinėje pusėje antros sekcijos footeryje
footer_p = new_section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_field(footer_p, "PAGE \\* MERGEFORMAT")
for r in footer_p.runs:
    set_run_font(r, size=FONT_SIZE_BODY)

# Pradėti numeravimą nuo 2 (kadangi titulinis = 1, bet jame nematomas)
sectPr = new_section._sectPr
pgNumType = OxmlElement('w:pgNumType')
pgNumType.set(qn('w:start'), '2')
sectPr.append(pgNumType)


# ============================================
# 2. TURINYS (automatinis, naujame lape, 14pt didžiosiomis)
# ============================================
p_turinys = doc.add_paragraph()
set_paragraph_format(p_turinys, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                     line_spacing=1.5, first_line_indent=None,
                     space_before=0, space_after=12)
r = p_turinys.add_run("TURINYS")
set_run_font(r, size=FONT_SIZE_H1, bold=True)
# „TURINYS" stiliai netaikomi — laikomės reikalavimo

# Automatinis turinys (Word atnaujins atidarius, paspaudus F9)
p_toc = doc.add_paragraph()
set_paragraph_format(p_toc, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                     line_spacing=1.5, first_line_indent=None,
                     space_before=0, space_after=0)
add_field(p_toc, 'TOC \\o "1-3" \\h \\z \\u')

# Pranešimas vartotojui apie turinio atnaujinimą
p_note = doc.add_paragraph()
set_paragraph_format(p_note, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                     line_spacing=1.0, first_line_indent=None,
                     space_before=12, space_after=0)
r = p_note.add_run(
    "(Atidarius dokumentą Word programoje, paspauskite dešinį pelės klavišą "
    "ant turinio ir pasirinkite „Update Field\" arba „Atnaujinti lauką\", "
    "kad turinys būtų sugeneruotas automatiškai.)"
)
set_run_font(r, size=FONT_SIZE_SOURCE, italic=True)


# ============================================
# 3. ĮVADAS (nenumeruotas skyrius)
# ============================================
add_h1(doc, "Įvadas", numbered=False)

add_body_paragraph(doc,
    "Šių dienų skaitmeninėje aplinkoje informacinės sistemos (toliau – IS) ir "
    "duomenų bazės (toliau – DB) yra neatsiejami komponentai, užtikrinantys "
    "efektyvų organizacijų darbą ir teikiamų paslaugų kokybę. Nors abi technologijos "
    "yra glaudžiai susijusios, jų kūrimo reikalavimai skiriasi, atspindėdami "
    "skirtingas paskirtis, apimtį ir naudotojų grupes. Norint kurti kokybiškas "
    "sistemas, būtina suprasti, kuo skiriasi IS reikalavimai nuo DB reikalavimų ir "
    "kaip jie tarpusavyje sąveikauja."
)

add_body_paragraph(doc,
    "Šio darbo tikslas – pasinaudojant interneto pagalba palyginti informacinių "
    "sistemų ir duomenų bazių kūrimo reikalavimus pagal pagrindinius aspektus "
    "(taikymo sritį, funkcionalumą, duomenų valdymą, saugumą, naudotojus, kūrimo "
    "procesą), pateikti realius pavyzdžius iš Lietuvos rinkos bei aprašyti "
    "praktikoje naudojamus reikalavimų rinkimo šablonus."
)

add_body_paragraph(doc, "Darbo uždaviniai:")

# Uždavinių sąrašas — be įtraukos
def add_list_item(doc, text):
    p = doc.add_paragraph(style='List Number')
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         line_spacing=1.5, first_line_indent=None,
                         space_before=0, space_after=0)
    for r in p.runs:
        r.text = ''
    r = p.add_run(text)
    set_run_font(r, size=FONT_SIZE_BODY)


add_list_item(doc, "Apžvelgti IS ir DB sampratą bei pagrindinius kūrimo aspektus.")
add_list_item(doc, "Atlikti IS ir DB reikalavimų palyginamąją analizę pateikiant struktūrinę lentelę.")
add_list_item(doc, "Išanalizuoti realų pavyzdį – Lietuvos elektroninės prekybos sistemą Pigu.lt ir logistikos paslaugų teikėjo Omniva Lietuva duomenų bazę.")
add_list_item(doc, "Aprašyti IS ir DB reikalavimų rinkimo šablonus, naudojamus praktikoje.")

add_body_paragraph(doc,
    "Darbo metodai: literatūros (mokslinių straipsnių, internetinių šaltinių) "
    "analizė, palyginamoji analizė, atvejo (case study) analizė. Darbas susideda "
    "iš trijų pagrindinių skyrių: pirmajame pristatoma IS ir DB samprata, "
    "antrajame atliekamas reikalavimų palyginimas su realiais pavyzdžiais, "
    "trečiajame aptariami reikalavimų rinkimo šablonai."
)


# ============================================
# 4. 1 SKYRIUS — IS IR DB SAMPRATA
# ============================================
add_h1(doc, "1. Informacinių sistemų ir duomenų bazių samprata")

add_body_paragraph(doc,
    "Šiame skyriuje aptariami pagrindiniai informacinių sistemų ir duomenų "
    "bazių apibrėžimai, jų paskirtis bei tarpusavio sąveika. Suprasti šių "
    "sąvokų esmę yra būtina prieš atliekant detalų reikalavimų palyginimą, "
    "nes kiekvienos technologijos paskirtis lemia jos kūrimo reikalavimus."
)

# 1.1
add_h2(doc, "1.1. Informacinių sistemų apibrėžimas ir paskirtis")

add_body_paragraph(doc,
    "Informacinė sistema – tai organizuotas žmonių, procesų, programinės bei "
    "techninės įrangos rinkinys, skirtas duomenims rinkti, apdoroti, saugoti "
    "ir paskirstyti, siekiant palaikyti organizacijos veiklą, sprendimų "
    "priėmimą bei valdymą (Laudon ir Laudon, 2020). IS aprėpia ne vien "
    "technologinį komponentą, bet ir verslo logiką, vartotojų sąsajas bei "
    "organizacinius procesus."
)

add_body_paragraph(doc,
    "Pagrindinė IS paskirtis – paversti duomenis informacija, kuri padeda "
    "organizacijai pasiekti savo tikslus. Tipiniai IS pavyzdžiai – įmonės "
    "išteklių valdymo sistemos (ERP), klientų ryšių valdymo sistemos (CRM) "
    "bei elektroninės prekybos platformos. IS reikalavimai apima ne tik "
    "techninius aspektus, bet ir vartotojų patirtį, integraciją su kitomis "
    "sistemomis, atitikimą teisės aktams (pvz., BDAR)."
)

# 1.2
add_h2(doc, "1.2. Duomenų bazių apibrėžimas ir paskirtis")

add_body_paragraph(doc,
    "Duomenų bazė – tai struktūruotas duomenų rinkinys, skirtas efektyviam "
    "duomenų saugojimui, paieškai, atnaujinimui ir valdymui (Connolly ir "
    "Begg, 2015). DB veikia per duomenų bazių valdymo sistemą (DBVS), kuri "
    "užtikrina duomenų vientisumą, saugumą ir prieinamumą. Populiariausios "
    "DBVS yra MySQL, PostgreSQL, Microsoft SQL Server, Oracle ir MongoDB."
)

add_body_paragraph(doc,
    "DB yra pagrindinis IS duomenų saugojimo komponentas, tačiau ji gali "
    "egzistuoti ir kaip atskira technologinė priemonė. Pagrindinė DB "
    "paskirtis – patikimai saugoti duomenis ir užtikrinti operatyvią "
    "prieigą prie jų. DB reikalavimai dažniausiai apima duomenų schemos "
    "projektavimą, normalizaciją, saugumo bei našumo užtikrinimą."
)

add_body_paragraph(doc,
    "Apibendrinant galima teigti, kad IS ir DB yra glaudžiai susijusios "
    "technologijos: IS naudoja DB kaip duomenų saugyklą, o DB struktūra ir "
    "funkcionalumas priklauso nuo IS keliamų reikalavimų. Toliau pateikiama "
    "detalesnė šių reikalavimų palyginamoji analizė."
)


# ============================================
# 5. 2 SKYRIUS — IS IR DB REIKALAVIMŲ PALYGINIMAS
# ============================================
add_h1(doc, "2. IS ir DB reikalavimų palyginimas")

add_body_paragraph(doc,
    "Šiame skyriuje pateikiama struktūrinė palyginamoji analizės lentelė, "
    "kurioje IS ir DB reikalavimai gretinami pagal septynis pagrindinius "
    "aspektus. Po lentelės pateikiama Pigu.lt ir Omniva Lietuva atvejo "
    "analizė bei rašytinė palyginamoji analizė."
)

# 2.1 — Lentelė
add_h2(doc, "2.1. Palyginamosios analizės lentelė")

add_body_paragraph(doc,
    "Pirmoje lentelėje (žr. 1 lentelę) pateikiamas struktūrinis IS ir DB "
    "kūrimo reikalavimų palyginimas pagal aspektus, nurodytus užduoties "
    "metodiniuose nurodymuose."
)

add_table_caption(doc, "1 lentelė", "IS ir DB kūrimo reikalavimų palyginimas")

# Lentelės kūrimas
table_data = [
    ["Aspektas", "Informacinės sistemos (IS) reikalavimai", "Duomenų bazės (DB) reikalavimai"],
    ["Taikymo sritis, tikslas",
     "Verslo procesų automatizavimas, vartotojų sąveika, sprendimų priėmimo palaikymas, paslaugų teikimas.",
     "Struktūruotas duomenų saugojimas, greita paieška, duomenų vientisumo ir prieinamumo užtikrinimas."],
    ["Funkcionalumas",
     "Verslo procesams skirtos funkcijos (užsakymai, ataskaitos, pranešimai), grafinė vartotojo sąsaja, integracijos su kitomis sistemomis per API.",
     "CRUD operacijos (kūrimas, skaitymas, atnaujinimas, šalinimas), užklausos (SQL), indeksavimas, transakcijų valdymas."],
    ["Duomenų valdymas",
     "Duomenų srautai tarp sistemos modulių, duomenų transformacija ir verslo logikos taikymas, duomenų pateikimas naudotojui.",
     "Duomenų schemos projektavimas, normalizacija, ryšiai tarp lentelių, ACID savybių užtikrinimas."],
    ["Saugumas",
     "Vartotojų autentifikacija (slaptažodis, dviejų faktorių), autorizacija pagal vaidmenis, sesijų valdymas, žurnalavimas, atitikimas BDAR.",
     "Šifravimas (saugomų ir perduodamų duomenų), prieigos teisės lentelėms ir laukams, atsarginės kopijos, audito sekos."],
    ["Vartotojai ir prieiga",
     "Galutiniai naudotojai (klientai, darbuotojai, administratoriai); prieiga per grafinę sąsają (web, mobilioji programėlė).",
     "DB administratoriai, programinės įrangos jungtys (API, ORM); prieiga per SQL klientą arba programinį kodą."],
    ["Kūrimo procesas",
     "SDLC etapai: reikalavimų rinkimas, analizė, projektavimas, kūrimas, testavimas, diegimas, priežiūra. Naudojamos Agile, Waterfall metodologijos.",
     "Konceptualus, loginis ir fizinis projektavimas (ER diagrama → reliacinė schema → DBVS); normalizacija, indeksų ir užklausų optimizavimas."],
    ["Realaus pasaulio pavyzdys",
     "Pigu.lt – Lietuvos elektroninės prekybos platforma, pateikianti vartotojo sąsają, prekių katalogą, krepšelį, mokėjimus ir klientų aptarnavimą.",
     "Omniva Lietuva siuntų sekimo DB – sauganti siuntų, klientų, paštomatų bei pristatymo būsenų informaciją, pasiekiama per API."],
]

table = doc.add_table(rows=len(table_data), cols=3)
table.style = 'Table Grid'
table.alignment = WD_ALIGN_PARAGRAPH.CENTER
table.autofit = False

# Lentelės plotis
col_widths = [Cm(3.8), Cm(6.3), Cm(6.3)]
for row in table.rows:
    for idx, cell in enumerate(row.cells):
        cell.width = col_widths[idx]

for row_idx, row_data in enumerate(table_data):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.rows[row_idx].cells[col_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        # Išvalykime egzistuojančią pastraipą
        cell.text = ""
        p = cell.paragraphs[0]
        set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                             line_spacing=1.0, first_line_indent=None,
                             space_before=2, space_after=2)
        r = p.add_run(cell_text)
        is_header = (row_idx == 0)
        set_run_font(r, size=FONT_SIZE_TABLE_DATA, bold=is_header)

# Šaltinis po lentele
add_source_note(doc, "Šaltinis: sudaryta autoriaus, remiantis Laudon ir Laudon (2020) bei Connolly ir Begg (2015).")


# 2.2 — Atvejo analizė
add_h2(doc, "2.2. Pigu.lt ir Omniva Lietuva pavyzdžių analizė")

add_body_paragraph(doc,
    "Pigu.lt yra viena didžiausių Lietuvos elektroninės prekybos platformų, "
    "jungianti šimtus tūkstančių prekių ir milijoną vartotojų. Šios "
    "informacinės sistemos kūrimo reikalavimai apima ne tik produktų "
    "katalogo valdymą, bet ir krepšelio funkcionalumą, mokėjimo apdorojimą "
    "(per Lietuvos bankus ir tarptautines mokėjimo sistemas), pristatymo "
    "logistiką, klientų aptarnavimą bei marketingo įrankius. Sistemos "
    "vartotojų sąsaja turi būti patogi tiek kompiuterio, tiek mobiliojo "
    "telefono ekrane."
)

add_body_paragraph(doc,
    "Omniva Lietuva, kaip pašto ir logistikos paslaugų teikėjas, valdo plačią "
    "paštomatų tinklo ir siuntų sekimo duomenų bazę. DB reikalavimai apima "
    "didžiulių duomenų kiekių apdorojimą realiu laiku, siuntų būsenų sekimą "
    "(kelionę nuo siuntėjo iki gavėjo), klientų asmens duomenų saugojimą "
    "pagal BDAR reikalavimus bei integraciją su partnerių sistemomis."
)

add_body_paragraph(doc,
    "Pigu.lt ir Omniva Lietuva sąveika rodo IS ir DB reikalavimų papildomumą "
    "praktikoje: kai pirkėjas atlieka užsakymą Pigu.lt platformoje ir "
    "pasirenka pristatymą per Omniva paštomatą, Pigu.lt informacinė sistema "
    "per API perduoda siuntimo duomenis Omniva sistemai, o pastaroji "
    "užregistruoja siuntą savo DB ir grąžina sekimo numerį atgal. Vartotojas "
    "tada gali sekti siuntą tiek Pigu.lt platformoje, tiek Omniva tinklalapyje "
    "ar mobiliojoje programėlėje."
)


# 2.3 — Rašytinė analizė (150–200 žodžių)
add_h2(doc, "2.3. Rašytinė palyginamoji analizė")

add_body_paragraph(doc,
    "Informacinės sistemos ir duomenų bazės yra technologiniai sprendimai, "
    "tačiau jų kūrimo reikalavimai skiriasi tiek tikslo, tiek detalumo "
    "aspektais. IS reikalavimai apima visą organizacinį kontekstą: vartotojų "
    "vaidmenis, verslo procesus, vartotojo sąveiką ir integracijas su kitomis "
    "sistemomis. Tuo tarpu DB reikalavimai sutelkti ties duomenų struktūra, "
    "jų vientisumu, saugumu ir efektyvia prieiga."
)

add_body_paragraph(doc,
    "Vis dėlto šios reikalavimų grupės papildo viena kitą. DB sukuriama "
    "remiantis IS funkciniais reikalavimais: nustačius, kokie verslo procesai "
    "bus automatizuoti, suprojektuojama, kokie duomenys turi būti saugomi. "
    "Pavyzdžiui, Pigu.lt elektroninės prekybos sistema reikalauja saugoti "
    "produktus, klientus, užsakymus ir mokėjimus – šie reikalavimai virsta "
    "konkrečia DB schema su atitinkamomis lentelėmis ir ryšiais."
)

add_body_paragraph(doc,
    "Realus pavyzdys – Pigu.lt sistema, sąveikaujanti su Omniva Lietuva DB "
    "per programines sąsajas: kai klientas užsako prekę, IS perduoda "
    "siuntimo duomenis Omniva sistemai, o ši saugo juos savo DB ir grąžina "
    "sekimo numerį atgal į IS. Toks ryšys parodo, kad IS reikalavimai "
    "apibrėžia, kokia informacija turi būti perduodama ir kaip ji "
    "pateikiama vartotojui, o DB reikalavimai nustato, kaip ji turi būti "
    "saugoma, struktūruojama ir apsaugota. Abi reikalavimų grupės yra "
    "vienodai svarbios kuriant patikimą skaitmeninę paslaugą."
)


# ============================================
# 6. 3 SKYRIUS — REIKALAVIMŲ RINKIMO ŠABLONAI
# ============================================
add_h1(doc, "3. IS ir DB reikalavimų rinkimo šablonai")

add_body_paragraph(doc,
    "Šiame skyriuje aprašomi dažniausiai naudojami šablonai informacinių "
    "sistemų ir duomenų bazių reikalavimams rinkti. Tinkamai parinkti "
    "šablonai padeda struktūrizuotai dokumentuoti reikalavimus, sumažina "
    "klaidų riziką ir padidina projekto sėkmės tikimybę."
)

# 3.1
add_h2(doc, "3.1. IS reikalavimų specifikacijos šablonai")

add_body_paragraph(doc,
    "Informacinių sistemų reikalavimams rinkti dažniausiai naudojami šie "
    "šablonai:"
)

# Bullet list
def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         line_spacing=1.5, first_line_indent=None,
                         space_before=0, space_after=0)
    for r in p.runs:
        r.text = ''
    r = p.add_run(text)
    set_run_font(r, size=FONT_SIZE_BODY)


add_bullet(doc,
    "SRS (angl. Software Requirements Specification) – pagal IEEE 830 "
    "standartą struktūrizuotas dokumentas, aprašantis funkcinius ir "
    "nefunkcinius sistemos reikalavimus, naudotojų klases, sąsajas ir "
    "veiklos apribojimus."
)
add_bullet(doc,
    "Panaudojimo atvejų (angl. Use Case) šablonas – aprašo vartotojų "
    "sąveiką su sistema, įskaitant aktorius, pagrindinį scenarijų bei "
    "alternatyvius eigos kelius. Dažnai pateikiamas UML diagramos forma."
)
add_bullet(doc,
    "Vartotojo istorijos (angl. User Stories) – Agile metodikoje "
    "naudojamas trumpas formatas: „Kaip [vartotojas], aš noriu "
    "[funkcionalumas], kad [tikslas]\". Tinka iteraciniam reikalavimų "
    "rinkimui."
)

# 3.2
add_h2(doc, "3.2. Duomenų bazių reikalavimų rinkimo šablonai")

add_body_paragraph(doc,
    "Duomenų bazių projektavimui ir reikalavimų dokumentavimui naudojami "
    "specializuoti šablonai:"
)

add_bullet(doc,
    "Esybių–ryšių diagrama (angl. Entity–Relationship Diagram, ERD) – "
    "grafinis modelis, vaizduojantis esybes (lenteles), atributus "
    "(laukus) ir ryšius tarp jų. Yra pagrindinis konceptualaus DB "
    "projektavimo įrankis."
)
add_bullet(doc,
    "Duomenų žodynas (angl. Data Dictionary) – lentelinis aprašymas, "
    "kuriame nurodomi visi DB elementai: lentelės, laukai, duomenų tipai, "
    "raktai, apribojimai ir laukų reikšmės."
)
add_bullet(doc,
    "Loginė ir fizinė schema – etapinio DB projektavimo dokumentai. "
    "Loginė schema aprašo lenteles, jų stulpelius ir ryšius nepriklausomai "
    "nuo DBVS, o fizinė schema – konkretų įgyvendinimą pasirinktoje DBVS."
)

# 3.3
add_h2(doc, "3.3. Šablonų reikšmė reikalavimų rinkimui")

add_body_paragraph(doc,
    "Standartizuoti šablonai padeda projektų komandoms tiksliai "
    "dokumentuoti reikalavimus, mažina nesusipratimų riziką tarp "
    "užsakovų, analitikų ir kūrėjų bei palengvina sistemos priežiūrą "
    "ir tobulinimą. Naudojant tokius įrankius kaip ERD ar SRS, "
    "projektai vystomi sistemingai – sumažėja klaidų skaičius, "
    "padidėja sistemos kokybė bei vartotojų pasitenkinimas."
)

add_body_paragraph(doc,
    "Apibendrinant galima teigti, kad IS ir DB reikalavimų rinkimo "
    "šablonai yra ne formalumas, o praktinė priemonė, padedanti pereiti "
    "nuo abstrakčių verslo poreikių prie konkretaus techninio sprendimo. "
    "Pasirinkti šablonai turi atitikti projekto mastą, naudojamą "
    "metodologiją bei komandos kompetenciją."
)


# ============================================
# 7. IŠVADOS
# ============================================
add_h1(doc, "Išvados", numbered=False)

add_list_item(doc,
    "Informacinės sistemos ir duomenų bazės yra glaudžiai susijusios, "
    "tačiau jų kūrimo reikalavimai skiriasi: IS aprėpia visą organizacijos "
    "veiklos kontekstą (verslo procesus, vartotojų sąveiką, integracijas), "
    "o DB sutelkta ties duomenų struktūra, vientisumu ir saugiu saugojimu."
)
add_list_item(doc,
    "IS reikalavimai apima funkcinius (verslo procesai, vartotojo sąveika) "
    "ir nefunkcinius aspektus (saugumas, našumas, atitiktis teisės aktams), "
    "o DB reikalavimai – duomenų schemą, normalizaciją, prieigos kontrolę "
    "bei ACID savybes."
)
add_list_item(doc,
    "Pigu.lt ir Omniva Lietuva pavyzdys parodo, kad IS ir DB reikalavimai "
    "praktikoje papildo vienas kitą: elektroninės prekybos platforma per "
    "API perduoda siuntimo duomenis logistikos paslaugų teikėjui, kuris "
    "juos saugo savo DB ir grąžina informaciją atgal į IS."
)
add_list_item(doc,
    "Tinkamai parinkti reikalavimų rinkimo šablonai (SRS, Use Case, User "
    "Stories, ERD, duomenų žodynas) padeda struktūrizuotai dokumentuoti "
    "reikalavimus, sumažina projektų klaidų riziką ir užtikrina aukštesnę "
    "kuriamos sistemos kokybę."
)


# ============================================
# 8. LITERATŪROS SĄRAŠAS (APA)
# ============================================
add_h1(doc, "Literatūros sąrašas", numbered=False)

literature = [
    "Connolly, T. ir Begg, C. (2015). Database Systems: A Practical Approach to Design, "
    "Implementation, and Management (6th ed.). Boston: Pearson Education.",

    "Laudon, K. C. ir Laudon, J. P. (2020). Management Information Systems: Managing the "
    "Digital Firm (16th ed.). Harlow: Pearson Education.",

    "IEEE Computer Society. (1998). IEEE Recommended Practice for Software Requirements "
    "Specifications (IEEE Std 830-1998). New York: Institute of Electrical and Electronics "
    "Engineers.",

    "Sommerville, I. (2016). Software Engineering (10th ed.). Boston: Pearson.",

    "Cohn, M. (2004). User Stories Applied: For Agile Software Development. Boston: "
    "Addison-Wesley Professional.",

    "Pigu.lt. (n.d.). Apie mus. Prieiga per internetą: https://pigu.lt/lt/apie-mus",

    "Omniva. (n.d.). Apie Omniva Lietuvoje. Prieiga per internetą: "
    "https://www.omniva.lt/privatus/apie-omniva",

    "Valstybinė duomenų apsaugos inspekcija. (2018). Bendrasis duomenų apsaugos reglamentas "
    "(BDAR). Prieiga per internetą: https://vdai.lrv.lt",
]

# APA stilius — pakabintos eilutės (hanging indent), atskira pastraipa kiekvienam šaltiniui
for src in literature:
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         line_spacing=1.5, first_line_indent=None,
                         space_before=0, space_after=6)
    # Pakabinta įtrauka (hanging) — pirmoji eilutė kairėje, kitos eilutės įtrauktos
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-1.25)
    r = p.add_run(src)
    set_run_font(r, size=FONT_SIZE_BODY)


# ============================================
# IŠSAUGOJIMAS
# ============================================
output_path = "/projects/sandbox/PU1_darbas/PU1_Andrius_Vargonas.docx"
doc.save(output_path)
print(f"Sukurta: {output_path}")
