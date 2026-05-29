"""
PU5 ataskaitos generatorius (.docx)

Tema: Duomenu bazes lenteliu kurimas pagal ERD
Autorius: Andrius Vargonas
"""

import os
from docx import Document
from docx.shared import Pt, Mm, Cm, RGBColor
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
    WD_TAB_ALIGNMENT, WD_TAB_LEADER,
)
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = "Times New Roman"
FIRST_LINE_INDENT = Cm(1.25)
HANGING_INDENT = Cm(1.25)
LINE_SPACING = 1.5
OUTPUT_FILE = "PU5_Vargonas.docx"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


def set_run_font(run, size=12, bold=False, italic=False, font_name=None):
    if font_name is None:
        font_name = FONT_NAME
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)


def set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         first_line_indent=None, left_indent=None,
                         line_spacing=LINE_SPACING,
                         space_before=0, space_after=0,
                         page_break_before=False):
    pf = p.paragraph_format
    pf.alignment = alignment
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if page_break_before:
        pf.page_break_before = True


def add_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  indent=True, size=12, bold=False, italic=False,
                  line_spacing=LINE_SPACING,
                  space_before=0, space_after=0,
                  page_break_before=False):
    p = doc.add_paragraph()
    set_paragraph_format(
        p, alignment=alignment,
        first_line_indent=(FIRST_LINE_INDENT if indent else Cm(0)),
        line_spacing=line_spacing,
        space_before=space_before, space_after=space_after,
        page_break_before=page_break_before,
    )
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_empty_line(doc, size=12):
    return add_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         indent=False, size=size)


def apply_heading_style(p, level):
    pPr = p._element.get_or_add_pPr()
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)
    pStyle.set(qn("w:val"), f"Heading{level}")


def add_heading_1(doc, text, page_break_before=True):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0), line_spacing=LINE_SPACING,
                         space_after=12, page_break_before=page_break_before)
    run = p.add_run(text.upper())
    set_run_font(run, size=14, bold=True)
    apply_heading_style(p, 1)
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                         first_line_indent=Cm(0), line_spacing=LINE_SPACING,
                         space_before=24, space_after=12)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True)
    apply_heading_style(p, 2)
    return p


def add_numbered_item(doc, number, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         first_line_indent=Cm(0), left_indent=Cm(0.5),
                         line_spacing=LINE_SPACING)
    run = p.add_run(f"{number}. {text}")
    set_run_font(run, size=12)
    return p


def add_bullet_item(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                         first_line_indent=Cm(0), left_indent=Cm(0.75),
                         line_spacing=LINE_SPACING)
    run = p.add_run("- " + text)
    set_run_font(run, size=12)
    return p


def add_table_caption(doc, number_text, title_text):
    p_num = doc.add_paragraph()
    set_paragraph_format(p_num, alignment=WD_ALIGN_PARAGRAPH.RIGHT,
                         first_line_indent=Cm(0),
                         line_spacing=LINE_SPACING, space_before=12)
    run = p_num.add_run(number_text)
    set_run_font(run, size=11, bold=True)

    p_title = doc.add_paragraph()
    set_paragraph_format(p_title, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0),
                         line_spacing=LINE_SPACING, space_after=6)
    run = p_title.add_run(title_text)
    set_run_font(run, size=11, bold=True, italic=True)


def add_figure_caption(doc, number_text, title_text):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0),
                         line_spacing=LINE_SPACING, space_before=4)
    run = p.add_run(f"{number_text}. {title_text}")
    set_run_font(run, size=11, bold=True)


def add_source_note(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0),
                         line_spacing=1.0, space_before=0, space_after=12)
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True)


def add_image(doc, filename, width_cm=15.0):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0),
                         line_spacing=1.0, space_before=12)
    run = p.add_run()
    img_path = os.path.join(SCRIPT_DIR, filename)
    if os.path.exists(img_path):
        run.add_picture(img_path, width=Cm(width_cm))
    else:
        run = p.add_run(f"[Paveikslo failas nerastas: {filename}]")
        set_run_font(run, size=10, italic=True)


def style_table_borders(table):
    tbl = table._element
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{name}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        tblBorders.append(b)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tblBorders)


def add_data_table(doc, headers, rows, col_widths_cm=None, data_size=10):
    n_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for cell in table.columns[i].cells:
                cell.width = Cm(w)
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                             first_line_indent=Cm(0), line_spacing=1.0)
        run = p.add_run(h)
        set_run_font(run, size=data_size, bold=True)
    for r_idx, row_data in enumerate(rows, start=1):
        row = table.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_format(p, alignment=align,
                                 first_line_indent=Cm(0), line_spacing=1.0)
            run = p.add_run(str(val))
            set_run_font(run, size=data_size)
    style_table_borders(table)
    return table


def add_bibliography_entry(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = LINE_SPACING
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.left_indent = HANGING_INDENT
    pf.first_line_indent = -HANGING_INDENT
    pf.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=12)


def add_appendix_header(doc, number_text, title_text):
    p_num = doc.add_paragraph()
    set_paragraph_format(p_num, alignment=WD_ALIGN_PARAGRAPH.RIGHT,
                         first_line_indent=Cm(0),
                         line_spacing=LINE_SPACING,
                         page_break_before=True)
    run = p_num.add_run(number_text.upper())
    set_run_font(run, size=12, bold=True)
    p_title = doc.add_paragraph()
    set_paragraph_format(p_title, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0),
                         line_spacing=LINE_SPACING,
                         space_before=12, space_after=18)
    run = p_title.add_run(title_text.upper())
    set_run_font(run, size=12, bold=True)


def setup_page(section):
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(15)
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(12.5)


def add_page_number_field(paragraph, size=12):
    run = paragraph.add_run()
    set_run_font(run, size=size)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE \\* MERGEFORMAT "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._element
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_separate)
    r.append(cached)
    r.append(fld_end)


def configure_section_with_page_numbers(section, first_page_no_number=False):
    setup_page(section)
    if first_page_no_number:
        section.different_first_page_header_footer = True
        first_footer = section.first_page_footer
        if first_footer.paragraphs:
            fp = first_footer.paragraphs[0]
        else:
            fp = first_footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer
    if footer.paragraphs:
        fp = footer.paragraphs[0]
    else:
        fp = footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_format(fp, alignment=WD_ALIGN_PARAGRAPH.RIGHT,
                         first_line_indent=Cm(0), line_spacing=1.0)
    add_page_number_field(fp, size=12)


def configure_default_style(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.space_after = Pt(0)
    for h_level, size, bold in [
        ("Heading 1", 14, True),
        ("Heading 2", 12, True),
        ("Heading 3", 12, True),
    ]:
        try:
            style = styles[h_level]
            style.font.name = FONT_NAME
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
        except KeyError:
            continue


def add_title_page(doc):
    add_paragraph(doc, "VILNIAUS UNIVERSITETAS",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=14, bold=True)
    add_paragraph(doc, "KAUNO FAKULTETAS",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=14, bold=True)
    add_empty_line(doc)
    add_paragraph(doc,
                  "SOCIALINIŲ MOKSLŲ IR TAIKOMOSIOS INFORMATIKOS INSTITUTAS",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=14)
    add_empty_line(doc)
    add_paragraph(doc, "Marketingo technologijų studijų programa",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=12)
    for _ in range(5):
        add_empty_line(doc)
    add_paragraph(doc, "ANDRIUS VARGONAS",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=12, bold=True)
    add_empty_line(doc)
    add_empty_line(doc)
    add_paragraph(doc, "DUOMENŲ BAZĖS LENTELIŲ KŪRIMAS PAGAL ERD",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=14, bold=True)
    add_empty_line(doc)
    add_paragraph(doc, "Praktinė užduotis Nr. 5",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=12, italic=True)
    add_empty_line(doc)
    add_paragraph(doc, "Informacijos sistemos ir duomenų bazės",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=12)
    for _ in range(7):
        add_empty_line(doc)
    add_paragraph(doc, "Kaunas",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=12)
    add_paragraph(doc, "2026",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER,
                  indent=False, size=12)


def add_toc_title(doc):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                         first_line_indent=Cm(0),
                         line_spacing=LINE_SPACING,
                         space_after=12, page_break_before=True)
    run = p.add_run("TURINYS")
    set_run_font(run, size=14, bold=True)


def add_toc_entry(doc, text, page_num, level=1, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = LINE_SPACING
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(0)
    if level == 2:
        pf.left_indent = Cm(0.75)
    pf.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run_text = p.add_run(text)
    set_run_font(run_text, size=12, bold=bold)
    run_num = p.add_run("\t" + str(page_num))
    set_run_font(run_num, size=12, bold=bold)


def add_toc(doc):
    add_toc_title(doc)
    add_toc_entry(doc, "ĮVADAS", 3, level=1, bold=True)
    add_toc_entry(doc, "1. ESYBIŲ-RYŠIŲ MODELIAVIMO PAGRINDAI", 4, level=1, bold=True)
    add_toc_entry(doc, "1.1. ER diagramos samprata ir paskirtis", 4, level=2)
    add_toc_entry(doc, "1.2. Esybės, atributai ir ryšiai", 5, level=2)
    add_toc_entry(doc, "1.3. Konvertavimas į reliacines lenteles", 6, level=2)
    add_toc_entry(doc, "2. TRIJŲ DALYKINIŲ SRIČIŲ ER DIAGRAMOS", 7, level=1, bold=True)
    add_toc_entry(doc, "2.1. E-prekybos sistema", 7, level=2)
    add_toc_entry(doc, "2.2. Ligoninės valdymo sistema", 9, level=2)
    add_toc_entry(doc, "2.3. Bibliotekos valdymo sistema", 11, level=2)
    add_toc_entry(doc, "3. DUOMENŲ BAZIŲ ĮGYVENDINIMAS", 13, level=1, bold=True)
    add_toc_entry(doc, "3.1. Pasirinktas DBVS įrankis", 13, level=2)
    add_toc_entry(doc, "3.2. E-prekybos duomenų bazė", 14, level=2)
    add_toc_entry(doc, "3.3. Ligoninės duomenų bazė", 15, level=2)
    add_toc_entry(doc, "3.4. Bibliotekos duomenų bazė", 16, level=2)
    add_toc_entry(doc, "3.5. Patirties aprašymas", 17, level=2)
    add_toc_entry(doc, "IŠVADOS", 18, level=1, bold=True)
    add_toc_entry(doc, "LITERATŪROS SĄRAŠAS", 19, level=1, bold=True)
    add_toc_entry(doc, "1 PRIEDAS. SQL DDL skriptai trims duomenų bazėms", 20, level=1, bold=True)
    add_toc_entry(doc, "11 PRIEDAS. Dirbtinio intelekto panaudojimo deklaracija", 22, level=1, bold=True)



# ============================================================
# IVADAS
# ============================================================

def add_introduction(doc):
    add_heading_1(doc, "ĮVADAS")

    add_paragraph(doc,
        "Šiuolaikinės informacinės sistemos remiasi tinkamai suprojektuotomis "
        "duomenų bazėmis, kuriose informacija saugoma struktūrizuotai. "
        "Esybių-ryšių diagrama (toliau - ER diagrama) yra tarp svarbiausių "
        "duomenų modeliavimo įrankių, leidžiančių aiškiai atvaizduoti "
        "dalykinės srities objektus, jų savybes ir tarpusavio ryšius "
        "(Chen, 1976; Connolly ir Begg, 2015). Marketingo technologijų "
        "studijų programos studentui ER diagramos samprata aktuali ne tik "
        "techniniu, bet ir verslo požiūriu - jos padeda susikalbėti su "
        "informacinių technologijų specialistais ir aiškiau formuluoti "
        "reikalavimus naujai kuriamoms sistemoms.")

    add_paragraph(doc,
        "Šio darbo tikslas - trims skirtingoms dalykinėms sritims sukurti ER "
        "diagramas ir, remiantis jomis, įgyvendinti tris reliacines duomenų "
        "bazes pasirinktoje DBVS aplinkoje.")

    add_paragraph(doc, "Darbo uždaviniai:", indent=True)
    add_numbered_item(doc, 1,
        "Apžvelgti ER modeliavimo pagrindus - sąvokas, žymėjimus ir "
        "konvertavimą į reliacines lenteles.")
    add_numbered_item(doc, 2,
        "Pateiktiems trims dalykinėms sritims (e-prekyba, ligoninė, biblioteka) "
        "atskirti esybes nuo atributų ir nustatyti tarpusavio ryšius.")
    add_numbered_item(doc, 3,
        "Sukurti tris ER diagramas, parodančias kiekvienos sistemos "
        "struktūrą.")
    add_numbered_item(doc, 4,
        "LibreOffice Base aplinkoje sukurti tris atskiras duomenų bazes su "
        "pirminiais ir išoriniais raktais, įvesti pavyzdinius duomenis ir "
        "trumpai aprašyti praktinę patirtį.")

    add_paragraph(doc,
        "Darbo metodai - mokslinės literatūros analizė, dalykinės srities "
        "modeliavimas, ER diagramų sudarymas, praktinis duomenų bazių "
        "kūrimas LibreOffice Base aplinkoje.")

    add_paragraph(doc,
        "Darbą sudaro trys skyriai. Pirmajame pateikiama ER modeliavimo "
        "teorinė apžvalga. Antrajame pristatomos trys dalykinės sritys ir "
        "joms sukurtos ER diagramos. Trečiajame skyriuje aprašomas "
        "praktinis trijų duomenų bazių sukūrimas LibreOffice Base aplinkoje.")

    # DI deklaracija (trumpa)
    add_paragraph(doc,
        "Rengiant šį darbą buvo naudotas dirbtinio intelekto įrankis "
        "Anthropic Claude (Sonnet 4.5, 2026 m. gegužės mėn.) tik teorinio "
        "teksto, esybių ir atributų paaiškinimų bei pavyzdinių duomenų "
        "informacijos generavimui. Praktinę darbo dalį - duomenų bazių "
        "kūrimą, ER diagramų projektavimą, lentelių struktūros sudarymą ir "
        "duomenų įvedimą LibreOffice Base aplinkoje - savarankiškai atliko "
        "darbo autorius. Detalesnis aprašymas pateiktas 11 priede.")


# ============================================================
# 1 SKYRIUS - ER pagrindai
# ============================================================

def add_chapter_1(doc):
    add_heading_1(doc, "1. ESYBIŲ-RYŠIŲ MODELIAVIMO PAGRINDAI")

    add_paragraph(doc,
        "Šiame skyriuje pateikiama esybių-ryšių modeliavimo teorinė apžvalga. "
        "Apibrėžiamos pagrindinės sąvokos, paaiškinama, kaip ER diagramos "
        "konvertuojamos į reliacines lenteles, ir kokie pagrindiniai "
        "principai užtikrina duomenų vientisumą.")

    add_heading_2(doc, "1.1. ER diagramos samprata ir paskirtis")
    add_paragraph(doc,
        "Esybių-ryšių diagramos modelį 1976 metais pasiūlė Peter Chen "
        "(Chen, 1976). Šis konceptualus duomenų modelis leidžia grafiškai "
        "atvaizduoti dalykinės srities esybes, jų atributus ir tarpusavio "
        "ryšius. ER diagrama padeda dar prieš kuriant duomenų bazę aiškiai "
        "suformuluoti, kokie objektai bus saugomi, kokios jų charakteristikos "
        "ir kaip jie tarpusavyje susiję.")
    add_paragraph(doc,
        "ER diagrama atlieka dvi pagrindines funkcijas. Pirma, ji yra "
        "komunikacinis įrankis tarp užsakovo ir kūrėjo - leidžia įsitikinti, "
        "kad abi pusės vienodai supranta sistemos paskirtį. Antra, ji "
        "tarnauja kaip techninis pagrindas, pagal kurį vėliau kuriamos "
        "reliacinės duomenų bazių lentelės.")

    add_heading_2(doc, "1.2. Esybės, atributai ir ryšiai")
    add_paragraph(doc,
        "Esybė (angl. entity) yra dalykinės srities objektas, apie kurį "
        "norima saugoti informaciją. Tipiniai pavyzdžiai - klientas, "
        "produktas, užsakymas, pacientas, knyga. Kiekviena esybė turi "
        "atributus (angl. attributes) - savybes, apibūdinančias jos "
        "konkrečius pavyzdžius. Pavyzdžiui, esybės klientas atributai gali "
        "būti vardas, pavardė, el. pašto adresas, registracijos data.")
    add_paragraph(doc,
        "Ryšys (angl. relationship) susieja dvi ar daugiau esybių. Ryšio "
        "kardinalumas (angl. cardinality) nurodo, kiek vienos esybės "
        "egzempliorių gali būti susieta su kitos esybės egzemplioriais. "
        "Pagrindiniai kardinalumo tipai: vienas-prie-vieno (1:1), "
        "vienas-prie-daug (1:N) ir daug-prie-daug (M:N). Šiame darbe "
        "pagrindinis modeliavimo principas yra 1:N ryšys, kuris natūraliai "
        "atsiranda tarp pagrindinių esybių (pvz., klientas-užsakymas) ir "
        "dažniausiai įgyvendinamas per išorinius raktus.")

    add_heading_2(doc, "1.3. Konvertavimas į reliacines lenteles")
    add_paragraph(doc,
        "ER diagrama yra konceptualus modelis - jis nurodo, ką saugoti, "
        "tačiau ne kaip tai daryti techniškai. Norint įgyvendinti diagramą "
        "reliacinėje duomenų bazėje, kiekviena esybė virsta lentele, o "
        "kiekvienas atributas - lentelės stulpeliu. Esybės identifikatorius "
        "(pvz., KlientoID) tampa pirminiu raktu (angl. primary key, PK), "
        "kuris unikaliai identifikuoja kiekvieną įrašą.")
    add_paragraph(doc,
        "Vienas-prie-daug (1:N) ryšys įgyvendinamas išoriniu raktu (angl. "
        "foreign key, FK). Daug pusės lentelėje pridedamas papildomas "
        "stulpelis, nukreipiantis į vienas pusės pirminį raktą. "
        "Pavyzdžiui, jei klientas turi daug užsakymų, lentelėje "
        "Uzsakymai yra stulpelis KlientoID, kuris nurodo konkretų "
        "lentelės Klientai įrašą. Toks sprendimas užtikrina referencinį "
        "duomenų vientisumą (Connolly ir Begg, 2015).")
    add_paragraph(doc,
        "Apibendrinant pirmąjį skyrių galima teigti, kad ER diagrama yra "
        "esminis tarpinis žingsnis tarp dalykinės srities supratimo ir "
        "techninio duomenų bazės įgyvendinimo. Tinkamai išskirtos esybės, "
        "atributai ir ryšiai sudaro pagrindą tolesniam darbui.")


# ============================================================
# 2 SKYRIUS - 3 ER diagramos
# ============================================================

def add_chapter_2(doc):
    add_heading_1(doc, "2. TRIJŲ DALYKINIŲ SRIČIŲ ER DIAGRAMOS")

    add_paragraph(doc,
        "Šiame skyriuje pristatomos trys dalykinės sritys, kurioms užduotyje "
        "pateikti elementų sąrašai. Kiekvienai sričiai atskirtos esybės nuo "
        "atributų, nustatyti ryšiai ir sukurta atitinkama ER diagrama.")

    # 2.1 E-prekyba
    add_heading_2(doc, "2.1. E-prekybos sistema")
    add_paragraph(doc,
        "Internetinės mažmeninės prekybos sistemos elementų sąraše pateikti "
        "klientai, užsakymai, produktai, taip pat įvairūs jų atributai. "
        "Atlikus analizę išskirtos trys pagrindinės esybės: Klientai, "
        "Produktai ir Užsakymai. Užsakymas yra siejantis objektas, kuris "
        "konkretų klientą susieja su konkrečiu produktu, pridedant užsakymo "
        "datą, kiekį ir pristatymo adresą.")

    add_paragraph(doc,
        "Esybių ir atributų atskyrimas pateiktas 1 lentelėje (žr. 1 lentelę).",
        indent=True)
    add_table_caption(doc, "1 lentelė",
                      "E-prekybos sistemos esybės ir atributai")
    headers = ["Esybė", "Atributai", "Tipas"]
    rows = [
        ["Klientai", "KlientoID, Vardas, ElPastas",
         "Pagrindinė esybė"],
        ["Produktai", "ProduktoID, Pavadinimas, Kaina, Kategorija",
         "Pagrindinė esybė"],
        ["Užsakymai",
         "UzsakymoID, KlientoID (FK), ProduktoID (FK), UzsakymoData, Kiekis, PristatymoAdresas",
         "Siejanti esybė"],
    ]
    add_data_table(doc, headers, rows, col_widths_cm=[2.8, 7.5, 5.0])
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    add_paragraph(doc,
        "Sistemos ER diagrama pateikta 1 paveiksle (žr. 1 pav.). Tarp "
        "Klientai ir Užsakymai, taip pat tarp Produktai ir Užsakymai yra "
        "1:N ryšys - vienas klientas gali turėti kelis užsakymus, o vienas "
        "produktas gali būti įtrauktas į kelis užsakymus.")
    add_image(doc, "img1_er_eprekyba.png", width_cm=15.5)
    add_figure_caption(doc, "1 pav",
                       "E-prekybos sistemos esybių-ryšių diagrama")
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    # 2.2 Ligonine
    add_heading_2(doc, "2.2. Ligoninės valdymo sistema")
    add_paragraph(doc,
        "Ligoninės valdymo sistemos elementų sąraše išskirtos trys "
        "pagrindinės esybės: Pacientai, Gydytojai ir Vizitai. Vizitas yra "
        "siejanti esybė, kuri konkretų pacientą susieja su konkrečiu "
        "gydytoju, pridedant vizito datą, diagnozę ir gydymo pastabas.")

    add_table_caption(doc, "2 lentelė",
                      "Ligoninės valdymo sistemos esybės ir atributai")
    headers = ["Esybė", "Atributai", "Tipas"]
    rows = [
        ["Pacientai", "PacientoID, PacientoVardas",
         "Pagrindinė esybė"],
        ["Gydytojai", "GydytojoID, GydytojoVardas, Specializacija",
         "Pagrindinė esybė"],
        ["Vizitai",
         "VizitoID, PacientoID (FK), GydytojoID (FK), VizitoData, Diagnoze, GydymoPastabos",
         "Siejanti esybė"],
    ]
    add_data_table(doc, headers, rows, col_widths_cm=[2.8, 7.5, 5.0])
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    add_paragraph(doc,
        "Sistemos ER diagrama pateikta 2 paveiksle (žr. 2 pav.). Vienas "
        "pacientas per laiką gali turėti kelis vizitus, o vienas gydytojas "
        "priima daug pacientų - tai natūralus 1:N ryšys.")
    add_image(doc, "img2_er_ligonine.png", width_cm=15.5)
    add_figure_caption(doc, "2 pav",
                       "Ligoninės valdymo sistemos esybių-ryšių diagrama")
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    # 2.3 Biblioteka
    add_heading_2(doc, "2.3. Bibliotekos valdymo sistema")
    add_paragraph(doc,
        "Bibliotekos valdymo sistemos elementų sąraše išskirtos trys "
        "pagrindinės esybės: Knygos, Nariai ir Skolinimasi. Skolinimosi "
        "įrašas yra siejanti esybė, kuri konkrečią knygą susieja su "
        "konkrečiu nariu, pridedant skolinimosi ir grąžinimo datas.")

    add_table_caption(doc, "3 lentelė",
                      "Bibliotekos valdymo sistemos esybės ir atributai")
    headers = ["Esybė", "Atributai", "Tipas"]
    rows = [
        ["Knygos", "KnygosID, Pavadinimas, Autorius, ISBN",
         "Pagrindinė esybė"],
        ["Nariai", "NarioID, NarioVardas, NarystesData",
         "Pagrindinė esybė"],
        ["Skolinimasi",
         "SkolinimoID, KnygosID (FK), NarioID (FK), SkolinimosiData, GrazinimoData",
         "Siejanti esybė"],
    ]
    add_data_table(doc, headers, rows, col_widths_cm=[2.8, 7.5, 5.0])
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    add_paragraph(doc,
        "Sistemos ER diagrama pateikta 3 paveiksle (žr. 3 pav.). Viena "
        "knyga per laiką gali būti pasiskolinta kelis kartus, o vienas "
        "narys gali pasiskolinti daug knygų. Toks dvigubas 1:N ryšys yra "
        "klasikinė bibliotekos valdymo modelio dalis.")
    add_image(doc, "img3_er_biblioteka.png", width_cm=15.5)
    add_figure_caption(doc, "3 pav",
                       "Bibliotekos valdymo sistemos esybių-ryšių diagrama")
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    add_paragraph(doc,
        "Apibendrinant antrąjį skyrių galima teigti, kad visos trys "
        "dalykinės sritys remiasi ta pačia struktūrine schema - dvi "
        "pagrindinės esybės ir viena siejanti esybė tarp jų. Toks modelis "
        "yra dažnas reliacinių duomenų bazių praktikoje ir tinka daugumai "
        "verslo informacinių sistemų.")



# ============================================================
# 3 SKYRIUS - DB igyvendinimas
# ============================================================

def add_chapter_3(doc):
    add_heading_1(doc, "3. DUOMENŲ BAZIŲ ĮGYVENDINIMAS")

    add_paragraph(doc,
        "Šiame skyriuje aprašomas trijų duomenų bazių praktinis sukūrimas "
        "LibreOffice Base aplinkoje pagal antrajame skyriuje sudarytas ER "
        "diagramas.")

    add_heading_2(doc, "3.1. Pasirinktas DBVS įrankis")
    add_paragraph(doc,
        "Praktinei užduoties daliai pasirinktas LibreOffice Base - "
        "atvirojo kodo darbalaukio duomenų bazės valdymo įrankis, "
        "priklausantis LibreOffice biuro paketui (The Document Foundation, "
        "2026). Pasirinkimą lėmė trys priežastys: nemokama licencija, "
        "kelių platformų palaikymas (Windows, macOS, Linux) ir tinkamumas "
        "akademinei aplinkai. Įrankis naudoja embedded HSQLDB variklį, "
        "kuris automatiškai paleidžiamas atidarant .odb failą ir "
        "užtikrina pilną reliacinį funkcionalumą - pirminius raktus, "
        "išorinius raktus ir referencinį vientisumą.")

    # 3.2 e-prekyba
    add_heading_2(doc, "3.2. E-prekybos duomenų bazė")
    add_paragraph(doc,
        "Pirmoji duomenų bazė sukurta failu eprekyba.odb. Joje suprojektuotos "
        "trys lentelės pagal 1 paveiksle pateiktą ER diagramą - Klientai, "
        "Produktai ir Užsakymai. Pirminiai raktai sukurti su AutoValue "
        "savybe, kad LibreOffice Base automatiškai priskirtų sekantį numerį "
        "kiekvienam naujam įrašui. Užsakymai lentelės du išoriniai raktai "
        "(KlientoID, ProduktoID) sukurti per Tools - Relationships langą, "
        "užtikrinant duomenų vientisumą tarp lentelių.")
    add_paragraph(doc,
        "Į kiekvieną lentelę įvesta po 5 įrašus, atspindinčius realių "
        "elektroninės parduotuvės klientų, produktų ir užsakymų pavyzdžius. "
        "Užsakymų lentelės datasheet view su 5 įrašais pavaizduotas "
        "4 paveiksle (žr. 4 pav.). Pilnas SQL DDL skriptas pateiktas "
        "1 priede.")
    add_image(doc, "img4_data_eprekyba.png", width_cm=15.5)
    add_figure_caption(doc, "4 pav",
                       "E-prekybos duomenų bazė: Užsakymai lentelė su 5 įrašais")
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    # 3.3 ligonine
    add_heading_2(doc, "3.3. Ligoninės duomenų bazė")
    add_paragraph(doc,
        "Antroji duomenų bazė sukurta failu ligonine.odb. Joje suprojektuotos "
        "trys lentelės pagal 2 paveiksle pateiktą ER diagramą - Pacientai, "
        "Gydytojai ir Vizitai. Pirminiai raktai sukurti su AutoValue "
        "savybe. Vizitai lentelės du išoriniai raktai (PacientoID, "
        "GydytojoID) užtikrina, kad nebūtų galima įvesti vizito su "
        "neegzistuojančiu pacientu ar gydytoju.")
    add_paragraph(doc,
        "Į kiekvieną lentelę įvesta po 5 įrašus, atspindinčius realią "
        "ligoninės situaciją - pacientai su vardais, gydytojai su "
        "specializacijomis ir vizitai su diagnozėmis bei gydymo pastabomis. "
        "Vizitų lentelės datasheet view pavaizduotas 5 paveiksle "
        "(žr. 5 pav.).")
    add_image(doc, "img5_data_ligonine.png", width_cm=15.5)
    add_figure_caption(doc, "5 pav",
                       "Ligoninės duomenų bazė: Vizitai lentelė su 5 įrašais")
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    # 3.4 biblioteka
    add_heading_2(doc, "3.4. Bibliotekos duomenų bazė")
    add_paragraph(doc,
        "Trečioji duomenų bazė sukurta failu biblioteka.odb. Joje "
        "suprojektuotos trys lentelės pagal 3 paveiksle pateiktą ER "
        "diagramą - Knygos, Nariai ir Skolinimasi. Knygų lentelėje "
        "ISBN laukui pridėtas UNIQUE apribojimas, kad nebūtų galima įvesti "
        "tos pačios knygos du kartus. Skolinimosi lentelėje du išoriniai "
        "raktai (KnygosID, NarioID) susieja su pagrindinėmis lentelėmis.")
    add_paragraph(doc,
        "Į kiekvieną lentelę įvesta po 5 įrašus - knygos su autoriais ir "
        "ISBN kodais, nariai su narystės datomis, skolinimosi įrašai su "
        "datomis. Skolinimosi lentelės datasheet view pavaizduotas "
        "6 paveiksle (žr. 6 pav.).")
    add_image(doc, "img6_data_biblioteka.png", width_cm=15.5)
    add_figure_caption(doc, "6 pav",
                       "Bibliotekos duomenų bazė: Skolinimasi lentelė su 5 įrašais")
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")

    # 3.5 Patirtis
    add_heading_2(doc, "3.5. Patirties aprašymas")
    add_paragraph(doc,
        "Praktinis darbas su LibreOffice Base aplinka parodė, kad įrankis "
        "yra pakankamai galingas akademinei ir smulkiojo verslo praktikai. "
        "Lentelių kūrimas per Table Design rėžimą yra intuityvus - "
        "kiekvienam laukui galima nustatyti tipą, ilgį, AutoValue ir kitas "
        "savybes. Pirminiai raktai pažymimi raktelio simboliu, o išoriniai "
        "raktai sukuriami per Tools - Relationships langą, kuriame "
        "vizualiai matomi visų lentelių ryšiai. Trijų skirtingų duomenų "
        "bazių sukūrimas užtruko apie dvi valandas, o sukurti modeliai "
        "pasižymi referenciniu vientisumu - bandymas įvesti neegzistuojantį "
        "išorinio rakto ID grąžina aiškų klaidos pranešimą.")

    add_paragraph(doc,
        "Apibendrinant trečiąjį skyrių galima teigti, kad LibreOffice Base "
        "puikiai tinka šio tipo akademinėms užduotims. Trijų skirtingų "
        "dalykinių sričių duomenų bazių sukūrimas leido praktiškai "
        "išbandyti pirminių ir išorinių raktų panaudojimą bei įsitikinti, "
        "kad ER diagrama tikrai yra naudingas projektavimo įrankis.")


# ============================================================
# ISVADOS
# ============================================================

def add_conclusions(doc):
    add_heading_1(doc, "IŠVADOS")
    add_numbered_item(doc, 1,
        "Esybių-ryšių diagrama yra esminis tarpinis žingsnis tarp dalykinės "
        "srities supratimo ir reliacinės duomenų bazės įgyvendinimo. ER "
        "modelis padeda aiškiai atskirti esybes nuo atributų bei nustatyti "
        "tarpusavio ryšius dar prieš kuriant lenteles.")
    add_numbered_item(doc, 2,
        "Trims dalykinėms sritims (e-prekyba, ligoninė, biblioteka) "
        "išskirtos po tris esybes - dvi pagrindinės ir viena siejanti. "
        "Toks struktūrinis šablonas atspindi klasikinį 1:N santykį tarp "
        "pagrindinių esybių, įgyvendinamą per išorinius raktus.")
    add_numbered_item(doc, 3,
        "Trys duomenų bazės (eprekyba.odb, ligonine.odb, biblioteka.odb) "
        "sėkmingai sukurtos LibreOffice Base aplinkoje. Kiekvienoje "
        "lentelėje įvesta po 5 įrašus, sukurti pirminiai ir išoriniai "
        "raktai užtikrina referencinį vientisumą.")
    add_numbered_item(doc, 4,
        "LibreOffice Base patogus mokomajai praktikai - jo Table Design "
        "rėžimas leidžia greitai sukurti lenteles, o Relationships langas "
        "vizualizuoja išorinius raktus. Įrankis tinka tiek studijoms, tiek "
        "smulkiems verslo sprendimams.")


# ============================================================
# LITERATUROS SARASAS
# ============================================================

def add_bibliography(doc):
    add_heading_1(doc, "LITERATŪROS SĄRAŠAS")
    entries = [
        "Chen, P. P. (1976). The Entity-Relationship Model - Toward a "
        "Unified View of Data. ACM Transactions on Database Systems, "
        "1(1), 9-36.",

        "Connolly, T. ir Begg, C. (2015). Database Systems: A Practical "
        "Approach to Design, Implementation, and Management (6th ed.). "
        "Boston: Pearson.",

        "Elmasri, R. ir Navathe, S. B. (2016). Fundamentals of Database "
        "Systems (7th ed.). Boston: Pearson.",

        "Hoffer, J. A., Ramesh, V. ir Topi, H. (2019). Modern Database "
        "Management (13th ed.). Boston: Pearson.",

        "The Document Foundation. (2026). LibreOffice Base Handbook. "
        "Prieiga per internetą: https://documentation.libreoffice.org",

        "Vilniaus universitetas. (2024). Dirbtinio intelekto naudojimo "
        "gairės (Nr. SPN-54). Vilnius: VU.",
    ]
    for e in entries:
        add_bibliography_entry(doc, e)


# ============================================================
# 1 PRIEDAS - SQL DDL skriptai
# ============================================================

def add_appendix_1(doc):
    add_appendix_header(doc, "1 priedas",
                        "SQL DDL skriptai trims duomenų bazėms")

    add_paragraph(doc,
        "Šiame priede pateikiami visų trijų duomenų bazių lentelių sukūrimo "
        "SQL skriptai (HSQLDB sintaksė, suderinama su LibreOffice Base "
        "embedded varikliu).")

    sql_blocks = [
        ("E-prekyba: CREATE TABLE Klientai",
         "CREATE TABLE Klientai (\n"
         "    KlientoID INTEGER IDENTITY,\n"
         "    Vardas VARCHAR(100) NOT NULL,\n"
         "    ElPastas VARCHAR(100),\n"
         "    CONSTRAINT uq_klientai_email UNIQUE (ElPastas)\n"
         ");"),

        ("E-prekyba: CREATE TABLE Produktai",
         "CREATE TABLE Produktai (\n"
         "    ProduktoID INTEGER IDENTITY,\n"
         "    Pavadinimas VARCHAR(100) NOT NULL,\n"
         "    Kaina DECIMAL(8,2),\n"
         "    Kategorija VARCHAR(50)\n"
         ");"),

        ("E-prekyba: CREATE TABLE Uzsakymai",
         "CREATE TABLE Uzsakymai (\n"
         "    UzsakymoID INTEGER IDENTITY,\n"
         "    KlientoID INTEGER NOT NULL,\n"
         "    ProduktoID INTEGER NOT NULL,\n"
         "    UzsakymoData DATE,\n"
         "    Kiekis INTEGER,\n"
         "    PristatymoAdresas VARCHAR(150),\n"
         "    FOREIGN KEY (KlientoID) REFERENCES Klientai(KlientoID),\n"
         "    FOREIGN KEY (ProduktoID) REFERENCES Produktai(ProduktoID)\n"
         ");"),

        ("Ligonine: CREATE TABLE Pacientai",
         "CREATE TABLE Pacientai (\n"
         "    PacientoID INTEGER IDENTITY,\n"
         "    PacientoVardas VARCHAR(100) NOT NULL\n"
         ");"),

        ("Ligonine: CREATE TABLE Gydytojai",
         "CREATE TABLE Gydytojai (\n"
         "    GydytojoID INTEGER IDENTITY,\n"
         "    GydytojoVardas VARCHAR(100) NOT NULL,\n"
         "    Specializacija VARCHAR(80)\n"
         ");"),

        ("Ligonine: CREATE TABLE Vizitai",
         "CREATE TABLE Vizitai (\n"
         "    VizitoID INTEGER IDENTITY,\n"
         "    PacientoID INTEGER NOT NULL,\n"
         "    GydytojoID INTEGER NOT NULL,\n"
         "    VizitoData DATE,\n"
         "    Diagnoze VARCHAR(150),\n"
         "    GydymoPastabos VARCHAR(255),\n"
         "    FOREIGN KEY (PacientoID) REFERENCES Pacientai(PacientoID),\n"
         "    FOREIGN KEY (GydytojoID) REFERENCES Gydytojai(GydytojoID)\n"
         ");"),

        ("Biblioteka: CREATE TABLE Knygos",
         "CREATE TABLE Knygos (\n"
         "    KnygosID INTEGER IDENTITY,\n"
         "    Pavadinimas VARCHAR(150) NOT NULL,\n"
         "    Autorius VARCHAR(100),\n"
         "    ISBN VARCHAR(20),\n"
         "    CONSTRAINT uq_isbn UNIQUE (ISBN)\n"
         ");"),

        ("Biblioteka: CREATE TABLE Nariai",
         "CREATE TABLE Nariai (\n"
         "    NarioID INTEGER IDENTITY,\n"
         "    NarioVardas VARCHAR(100) NOT NULL,\n"
         "    NarystesData DATE\n"
         ");"),

        ("Biblioteka: CREATE TABLE Skolinimasi",
         "CREATE TABLE Skolinimasi (\n"
         "    SkolinimoID INTEGER IDENTITY,\n"
         "    KnygosID INTEGER NOT NULL,\n"
         "    NarioID INTEGER NOT NULL,\n"
         "    SkolinimosiData DATE,\n"
         "    GrazinimoData DATE,\n"
         "    FOREIGN KEY (KnygosID) REFERENCES Knygos(KnygosID),\n"
         "    FOREIGN KEY (NarioID) REFERENCES Nariai(NarioID)\n"
         ");"),
    ]

    for label, sql in sql_blocks:
        p_label = doc.add_paragraph()
        set_paragraph_format(p_label, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                             first_line_indent=Cm(0),
                             line_spacing=LINE_SPACING, space_before=10)
        run = p_label.add_run(label)
        set_run_font(run, size=12, bold=True)
        p_sql = doc.add_paragraph()
        set_paragraph_format(p_sql, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                             first_line_indent=Cm(0),
                             left_indent=Cm(0.5),
                             line_spacing=1.0)
        run = p_sql.add_run(sql)
        set_run_font(run, size=10, font_name="Courier New")


# ============================================================
# 11 PRIEDAS - DI deklaracija (trumpa)
# ============================================================

def add_appendix_11(doc):
    add_appendix_header(doc, "11 priedas",
                        "Dirbtinio intelekto panaudojimo deklaracija")
    add_paragraph(doc,
        "Rengiant šį darbą buvo naudotas generatyvinis dirbtinio intelekto "
        "(toliau - DI) įrankis Anthropic Claude (Sonnet 4.5, internetinė "
        "prieiga, naudota 2026 m. gegužės mėnesį). Praktinę darbo dalį - "
        "ER diagramų projektavimą, lentelių struktūros sudarymą, pirminių "
        "ir išorinių raktų sąsajų nustatymą bei duomenų įvedimą - "
        "savarankiškai atliko darbo autorius LibreOffice Base aplinkoje. "
        "DI buvo panaudotas tik aprašomojo teksto, teorinių sąvokų "
        "paaiškinimų ir pavyzdinių duomenų informacijos pateikimui.")

    add_paragraph(doc,
        "DI sugeneruotas tekstas nebuvo perkeltas į darbą be peržiūros - "
        "kiekvienas teiginys patikrintas ir reikšmingai redaguotas autoriaus. "
        "DI sugeneruoto turinio dalis darbe neviršija 15 procentų. "
        "Autorius susipažinęs su Vilniaus universiteto 2024 m. patvirtintomis "
        "DI naudojimo gairėmis (Nr. SPN-54) ir prisiima visišką atsakomybę "
        "už darbo turinį.")

    add_table_caption(doc, "4 lentelė", "DI naudojimo apimties suvestinė")
    headers = ["Rodiklis", "Reikšmė"]
    rows = [
        ["DI modelis", "Anthropic Claude Sonnet 4.5"],
        ["Naudojimo data", "2026 m. gegužės mėn."],
        ["DI naudojimo tikslas",
         "Aprašomojo teksto, sąvokų paaiškinimų, pavyzdinių duomenų pateikimas"],
        ["Kas atlikta savarankiškai (autoriaus)",
         "ER diagramos, duomenų bazės, lentelės, raktai, duomenų įvedimas"],
        ["DI sugeneruoto turinio dalis darbe", "mažiau nei 15 proc."],
    ]
    add_data_table(doc, headers, rows, col_widths_cm=[6.0, 10.5])
    add_source_note(doc, "Šaltinis: sudaryta autoriaus.")


def main():
    doc = Document()
    configure_default_style(doc)
    section = doc.sections[0]
    configure_section_with_page_numbers(section, first_page_no_number=True)

    add_title_page(doc)
    add_toc(doc)
    add_introduction(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_conclusions(doc)
    add_bibliography(doc)
    add_appendix_1(doc)
    add_appendix_11(doc)

    out_path = os.path.join(SCRIPT_DIR, OUTPUT_FILE)
    doc.save(out_path)
    print(f"Sukurta: {out_path}")


if __name__ == "__main__":
    main()
